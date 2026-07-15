#!/usr/bin/env python3
"""Generate the Pownal Security Schedule as a print-ready xlsx (landscape letter, 1 page)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

YELLOW = PatternFill("solid", fgColor="F6C445")
GREEN_HDR = PatternFill("solid", fgColor="93C47D")
GREEN = PatternFill("solid", fgColor="A9D08E")
BLUE_HDR = PatternFill("solid", fgColor="6D9EEB")
BLUE = PatternFill("solid", fgColor="CFE2F3")

thin = Side(style="thin", color="666666")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

DAYS = [
    ("August 11th  TUESDAY",   ["5pm - 10pm", "10pm - 3am", "3am - 8am (when volunteers arrive)"]),
    ("August 12th  WEDNESDAY", ["5pm - 10pm", "10pm - 3am", "3am - 8am (when volunteers arrive)"]),
    ("August 13th  THURSDAY",  ["5pm - 10pm", "10pm - 3am", "3am - 8am (when volunteers arrive)"]),
    ("August 14th  FRIDAY",    ["5pm - 10pm", "10pm - 3am", "3am - 8am (when volunteers arrive)"]),
    ("August 15th  SATURDAY",  ["5pm - 10pm", "10pm - 3am", "3am - 8am (when volunteers arrive)"]),
    ("August 16th  SUNDAY",    ["7am - 12pm", "12pm - 5pm", "5pm - 10pm", "10pm - 3am", "3am - 8am"]),
    ("August 17th  MONDAY",    ["7am - 12pm", "12pm - 5pm", "5pm - 10pm", "10pm - 3am", "3am - 8am (when volunteers arrive)"]),
]

wb = Workbook()
ws = wb.active
ws.title = "Security Schedule"

# Block column layout: 3 blocks of (shift, vol1, vol2), separated by spacer columns.
# Blocks start at columns A(1), E(5), I(9); spacers at D(4), H(8).
BLOCK_COLS = [1, 5, 9]
for start in BLOCK_COLS:
    ws.column_dimensions[get_column_letter(start)].width = 18      # shift
    ws.column_dimensions[get_column_letter(start + 1)].width = 13  # volunteer 1
    ws.column_dimensions[get_column_letter(start + 2)].width = 13  # volunteer 2
for spacer in (4, 8):
    ws.column_dimensions[get_column_letter(spacer)].width = 1.5

# Title
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
t = ws.cell(row=1, column=1, value="POWNAL SECURITY SCHEDULE")
t.font = Font(bold=True, size=20)
t.alignment = Alignment(horizontal="center", vertical="center")
for c in range(1, 12):
    ws.cell(row=1, column=c).fill = YELLOW
    ws.cell(row=1, column=c).border = BORDER
ws.row_dimensions[1].height = 40
ws.row_dimensions[2].height = 6

def draw_day(row, col, title, shifts):
    """Draw one day block starting at (row, col). Returns rows used."""
    # Date header across the 3 block columns
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
    h = ws.cell(row=row, column=col, value=title)
    h.font = Font(bold=True, size=12)
    h.alignment = Alignment(horizontal="left", vertical="center")
    for c in range(col, col + 3):
        ws.cell(row=row, column=c).fill = GREEN_HDR
        ws.cell(row=row, column=c).border = BORDER
    ws.row_dimensions[row].height = 22

    # Column labels
    lr = row + 1
    for offset, (label, fill, color) in enumerate([
        ("Shift", GREEN_HDR, "000000"),
        ("Volunteer 1", BLUE_HDR, "FFFFFF"),
        ("Volunteer 2", BLUE_HDR, "FFFFFF"),
    ]):
        cell = ws.cell(row=lr, column=col + offset, value=label)
        cell.font = Font(bold=True, size=10, color=color)
        cell.fill = fill
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center" if offset else "left", vertical="center")
    ws.row_dimensions[lr].height = 18

    # Shift rows with two blank fill-in cells each
    for i, shift in enumerate(shifts):
        r = lr + 1 + i
        s = ws.cell(row=r, column=col, value=shift)
        s.font = Font(bold=True, size=9.5)
        s.fill = GREEN
        s.border = BORDER
        s.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for offset in (1, 2):
            v = ws.cell(row=r, column=col + offset)
            v.fill = BLUE
            v.border = BORDER
        ws.row_dimensions[r].height = 42
    return 2 + len(shifts)

# Lay out days in a 3 x 3 grid (bands of 3 days)
row = 3
for band_start in range(0, len(DAYS), 3):
    band = DAYS[band_start:band_start + 3]
    used = 0
    for i, (title, shifts) in enumerate(band):
        used = max(used, draw_day(row, BLOCK_COLS[i], title, shifts))
    row += used
    ws.row_dimensions[row].height = 10  # spacer between bands
    row += 1

# Footnote (from the original sheet)
NOTE = ("Please note that shifts are organized by calendar day. Even if a shift extends "
        "into the early morning hours, it is recorded under the day it begins to maintain "
        "scheduling consistency.")
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
n = ws.cell(row=row, column=1, value=NOTE)
n.font = Font(italic=True, size=9)
n.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[row].height = 24
row += 1

# Print setup: one landscape letter page
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
ws.page_setup.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.4)
ws.print_area = f"A1:K{row - 1}"
ws.sheet_view.showGridLines = False

wb.save("/home/user/BRXSTUDIOS/Pownal Security Schedule.xlsx")
print(f"Saved. Last row: {row - 1}")
